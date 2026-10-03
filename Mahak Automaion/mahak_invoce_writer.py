# -*- coding: utf-8 -*-
"""
نوشتن خودکار فاکتور فروش در محک — بر اساس توالی واقعی کلیدهایی که از
روی ضبط‌های شما استخراج شد.

نصب (یک‌بار):
    pip install openpyxl pynput pygetwindow

اجرا:
    python mahak_invoice_writer.py

⚠️ قبل از هر اجرای واقعی، حتماً یک بار با «حالت آزمایشی» (که به‌جای
محک در Notepad تایپ می‌کند) تست کنید. در حین اجرای واقعی هم می‌توانید
هر لحظه با زدن Esc متوقفش کنید.
"""

import os
import re
import sys
import json
import time
import difflib
import subprocess
import datetime
import threading
import tkinter as tk
from tkinter import filedialog, messagebox, ttk

import openpyxl
from pynput.keyboard import Controller, Key, KeyCode, Listener
import pygetwindow as gw

NUMPAD_PLUS = KeyCode.from_vk(0x6B)  # کلید "+" روی صفحه‌کلید عددی (Numpad Add)

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
SETTINGS_FILE = os.path.join(BASE_DIR, "writer_settings.json")

kb = Controller()

DEFAULT_SETTINGS = {
    "columns": {"invoice_no": "", "customer": "", "date": "", "item": "", "qty": "", "price": ""},
    "default_customer_code": "",
    "match_threshold": 0.5,
    "window_title": "",
    "columns_set": False,
    "window_set": False,
    "speed_multiplier": 3.0,
}

abort_flag = {"stop": False}
SPEED = {"value": 3.0}


def scaled(base_seconds):
    return base_seconds * SPEED["value"]


# ---------------------------------------------------------------------
# تنظیمات
# ---------------------------------------------------------------------
def load_settings():
    if os.path.exists(SETTINGS_FILE):
        with open(SETTINGS_FILE, "r", encoding="utf-8") as f:
            data = json.load(f)
        s = json.loads(json.dumps(DEFAULT_SETTINGS))
        s.update(data)
        return s
    return json.loads(json.dumps(DEFAULT_SETTINGS))


def save_settings(s):
    with open(SETTINGS_FILE, "w", encoding="utf-8") as f:
        json.dump(s, f, ensure_ascii=False, indent=2)


# ---------------------------------------------------------------------
# کلید اضطراری توقف (Esc) — در پس‌زمینه گوش می‌دهد
# ---------------------------------------------------------------------
def start_panic_listener():
    def on_press(key):
        if key == Key.esc:
            abort_flag["stop"] = True

    listener = Listener(on_press=on_press)
    listener.daemon = True
    listener.start()
    return listener


def check_abort():
    if abort_flag["stop"]:
        raise KeyboardInterrupt("متوقف شد توسط کاربر (Esc)")


# ---------------------------------------------------------------------
# ابزارهای تایپ سطح‌پایین
# ---------------------------------------------------------------------
KEY_HOLD = 0.04  # مدت نگه داشتن هر کلید؛ برخی کنترل‌های قدیمی کلیدِ خیلی کوتاه را نمی‌بینند


def _tap(key_obj):
    kb.press(key_obj)
    time.sleep(KEY_HOLD)
    kb.release(key_obj)


def type_text(text, delay=0.035):
    """
    ارقام را مثل یک کاربر واقعی، با «کلید رقم» (Virtual-Key) می‌فرستد، نه کاراکتر
    یونیکد؛ چون فیلدهای ماسک‌شده‌ی محک (مثل تاریخ) کاراکتر یونیکد تزریقی را
    نمی‌پذیرند. ارقام فارسی/عربی هم به 0-9 نرمال می‌شوند.
    """
    for ch in str(text):
        check_abort()
        if ch.isdigit():
            _tap(KeyCode.from_vk(0x30 + int(ch)))
        else:
            kb.type(ch)
        time.sleep(scaled(delay))


def press_key(key_obj, times=1, gap=0.15):
    for _ in range(times):
        check_abort()
        _tap(key_obj)
        time.sleep(scaled(gap))


def press_plus(gap=0.35):
    check_abort()
    _tap(NUMPAD_PLUS)
    time.sleep(scaled(gap))


def enter_tab_pairs(pairs, gap=0.15):
    for _ in range(pairs):
        press_key(Key.enter, 1, gap)
        press_key(Key.tab, 1, gap)


# ---------------------------------------------------------------------
# تبدیل تاریخ به فرمت روز-ماه-سال۲رقمی (مطابق چیزی که تأیید کردید)
# ---------------------------------------------------------------------
def num_to_str(value):
    """تبدیل امن عدد/متن به رشته، بدون اعشار زائد (مثل 102001.0)."""
    if value is None:
        return ""
    if isinstance(value, float):
        return str(int(value)) if value == int(value) else str(value)
    return str(value).strip()


def format_date_ddmmyy(raw_value):
    if isinstance(raw_value, (datetime.date, datetime.datetime)):
        raise ValueError(
            "ستون تاریخ در فایل اکسل به‌صورت تاریخ واقعی اکسل (میلادی) ذخیره شده، "
            "نه متن شمسی. لطفاً فرمت آن ستون را در فایل خام به متن شمسی (مثل 1405/07/04) تغییر دهید."
        )
    s = str(raw_value).strip()
    digits_only = re.sub(r"\D", "", s)
    parts = [p for p in re.split(r"[/\-\.]", s) if p.strip() != ""]

    if len(parts) == 3:
        year_part = None
        for p in parts:
            if len(p) == 4 or int(p) > 31:
                year_part = p
                break
        if year_part is None:
            year_part, month_part, day_part = parts[0], parts[1], parts[2]
        else:
            others = [p for p in parts if p != year_part]
            month_part, day_part = others[0], others[1]
        day = day_part.zfill(2)
        month = month_part.zfill(2)
        year2 = year_part[-2:].zfill(2)
        return day + month + year2

    if len(digits_only) == 6:
        return digits_only

    raise ValueError(f"فرمت تاریخ «{raw_value}» قابل تشخیص نیست.")


# ---------------------------------------------------------------------
# لیست اشخاص (کد + نام) و تطبیق فازی برای پیدا کردن کد مشتری
# ---------------------------------------------------------------------
def load_customer_lookup(path):
    """
    فایل اکسل «لیست اشخاص» را می‌خواند و [(کد, نام), ...] برمی‌گرداند.
    ابتدا سعی می‌کند ستون کد/نام را از روی عنوان ستون‌ها پیدا کند؛ اگر
    عنوانی پیدا نشد (مثلاً چون هنگام کپی/پیست از محک، فقط ردیف‌های داده
    منتقل شده‌اند)، None برمی‌گرداند تا کاربر خودش ستون‌ها را انتخاب کند.
    """
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        raise RuntimeError("فایل خالی است.")

    headers = [str(c) if c is not None else "" for c in rows[0]]
    code_col = None
    name_col = None
    for i, h in enumerate(headers):
        if code_col is None and "کد" in h:
            code_col = i
        if name_col is None and ("نام" in h or "خانوادگی" in h):
            name_col = i

    if code_col is not None and name_col is not None:
        result = []
        for row in rows[1:]:
            code = row[code_col] if code_col < len(row) else None
            name = row[name_col] if name_col < len(row) else None
            if code in (None, "") or name in (None, ""):
                continue
            result.append((num_to_str(code), str(name).strip()))
        return result

    return None  # عنوان ستون پیدا نشد؛ فراخوان باید انتخاب دستی انجام دهد


def load_customer_lookup_manual(path, code_col, name_col, skip_first_row):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    result = []
    for row in rows[1:] if skip_first_row else rows:
        code = row[code_col] if code_col < len(row) else None
        name = row[name_col] if name_col < len(row) else None
        if code in (None, "") or name in (None, ""):
            continue
        result.append((num_to_str(code), str(name).strip()))
    return result


def pick_customer_columns_dialog(path, on_done):
    """
    پنجره‌ای برای انتخاب دستی ستون کد/نام، وقتی عنوان ستون‌ها پیدا نشود.
    on_done(customer_list) پس از تأیید کاربر صدا زده می‌شود؛ در صورت لغو، صدا زده نمی‌شود.
    """
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        messagebox.showerror("خطا", "فایل خالی است.")
        return

    n_cols = max(len(r) for r in rows[:10])
    col_labels = [f"ستون {i+1}" for i in range(n_cols)]

    top = tk.Toplevel()
    top.title("انتخاب دستی ستون‌های کد و نام")
    top.geometry("700x400")
    tk.Label(
        top,
        text="عنوان ستون‌ها در این فایل پیدا نشد. چند ردیف اول فایل را می‌بینید — "
             "بگویید کدام ستون «کد» و کدام «نام» است:",
        wraplength=680, justify="right",
    ).pack(pady=8)

    preview = tk.Text(top, width=90, height=8)
    preview.pack(padx=10)
    for r in rows[:6]:
        preview.insert(tk.END, "   |   ".join(str(c) if c is not None else "" for c in r) + "\n")
    preview.config(state="disabled")

    frm = tk.Frame(top)
    frm.pack(pady=10)
    tk.Label(frm, text="ستون کد:").grid(row=0, column=1, padx=5)
    code_var = tk.StringVar(value=col_labels[0])
    ttk.Combobox(frm, textvariable=code_var, values=col_labels, state="readonly", width=15).grid(row=0, column=0, padx=5)
    tk.Label(frm, text="ستون نام:").grid(row=1, column=1, padx=5)
    name_var = tk.StringVar(value=col_labels[1] if len(col_labels) > 1 else col_labels[0])
    ttk.Combobox(frm, textvariable=name_var, values=col_labels, state="readonly", width=15).grid(row=1, column=0, padx=5)

    skip_var = tk.BooleanVar(value=False)
    tk.Checkbutton(top, text="ردیف اول فایل هم داده است (عنوان ستون نیست)، آن را هم نادیده نگیر", variable=skip_var).pack()

    def confirm():
        code_idx = col_labels.index(code_var.get())
        name_idx = col_labels.index(name_var.get())
        skip_first = not skip_var.get()
        result = load_customer_lookup_manual(path, code_idx, name_idx, skip_first)
        if not result:
            messagebox.showerror("خطا", "با این انتخاب هیچ داده‌ای پیدا نشد.")
            return
        top.destroy()
        on_done(result)

    tk.Button(top, text="تأیید", command=confirm, width=20, bg="#5cb85c", fg="white").pack(pady=10)


def find_customer_code(raw_name, customer_list, threshold, default_code):
    """
    بهترین تطابق (بر اساس درصد شباهت متن) را در لیست اشخاص پیدا می‌کند.
    اگر بالاترین شباهت کمتر از threshold باشد، default_code برگردانده می‌شود.
    """
    raw_name = (raw_name or "").strip()
    if not raw_name:
        return default_code, None, 0.0

    best_code, best_name, best_score = None, None, 0.0
    for code, name in customer_list:
        score = difflib.SequenceMatcher(None, raw_name, name).ratio()
        if score > best_score:
            best_score = score
            best_code = code
            best_name = name

    if best_score >= threshold:
        return best_code, best_name, best_score
    return default_code, None, best_score


# ---------------------------------------------------------------------
# توالی واقعی وارد کردن فاکتور (بر اساس ضبط‌های تأییدشده)
# ---------------------------------------------------------------------
def write_header(customer_code, date_raw, first_invoice):
    """
    توالی بدون Enter، فقط Tab (بر اساس ضبط‌های تأییدشده‌ی جدید):
    F2 (فقط فاکتور اول) → کد مشتری → Tab, Tab, Tab
    → تاریخ (۶ رقم: روز-ماه-سال دو رقمی) → Tab
    """
    if first_invoice:
        press_key(Key.f2, 1, 0.8)       # فقط یک‌بار، برای باز کردن صفحه‌ی فاکتور فروش
    type_text(customer_code)            # بعد از F2 فوکوس روی فیلد کد خریدار است
    press_key(Key.tab, 3, 0.3)

    date_digits = format_date_ddmmyy(date_raw)
    type_text(date_digits, delay=0.06)

    press_key(Key.tab, 1, 0.35)


def write_item(item_code, qty, price, is_last=None):
    """
    + → کد کالا → Tab, Tab → مقدار → Tab → قیمت واحد → F3
    (بدون Enter؛ بین دو کالا هیچ کلید اضافه‌ای لازم نیست)
    """
    press_plus(0.5)
    type_text(item_code)
    press_key(Key.tab, 2, 0.3)
    type_text(qty, delay=0.05)
    press_key(Key.tab, 1, 0.25)
    type_text(price, delay=0.05)
    press_key(Key.f3, 1, 0.7)           # ذخیره و خروج از پنجره‌ی انتخاب کالا


def write_invoice(customer_code, date_raw, items, first_invoice, log_func=print):
    write_header(customer_code, date_raw, first_invoice)
    for i, it in enumerate(items):
        check_abort()
        is_last = i == len(items) - 1
        write_item(it["code"], it["qty"], it["price"], is_last)
    press_key(Key.f3, 1, 0.6)           # ذخیره و جدید — فاکتور را ذخیره و صفحه را برای فاکتور بعدی آماده می‌کند
    log_func(f"  فاکتور مشتری با کد «{customer_code}» و {len(items)} ردیف کالا ثبت شد.")


# ---------------------------------------------------------------------
# فعال کردن پنجره‌ی مقصد (محک یا Notepad در حالت آزمایشی)
# ---------------------------------------------------------------------
def activate_window(title_substring):
    wins = gw.getWindowsWithTitle(title_substring)
    if not wins:
        raise RuntimeError(f"پنجره‌ای با عنوان شامل «{title_substring}» پیدا نشد.")
    win = wins[0]
    try:
        if win.isMinimized:
            win.restore()
        win.activate()
    except Exception:
        pass
    time.sleep(scaled(0.6))


# ---------------------------------------------------------------------
# خواندن فایل خام و گروه‌بندی به فاکتورها
# ---------------------------------------------------------------------
def build_invoices(raw_path, settings, customer_list):
    cols = settings["columns"]
    wb = openpyxl.load_workbook(raw_path, data_only=True)
    ws = wb.active
    headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]

    def idx(name):
        try:
            return headers.index(name)
        except ValueError:
            return None

    i_inv = idx(cols.get("invoice_no", ""))
    i_cust = idx(cols["customer"])
    i_date = idx(cols["date"])
    i_item = idx(cols["item"])
    i_qty = idx(cols["qty"])
    i_price = idx(cols["price"])

    missing = [k for k, v in [("invoice_no", i_inv), ("customer", i_cust), ("date", i_date), ("item", i_item), ("qty", i_qty), ("price", i_price)] if v is None]
    if missing:
        raise RuntimeError("ستون‌های زیر در فایل پیدا نشد (مرحله‌ی ۱ را دوباره انجام دهید): " + ", ".join(missing))

    default_code = settings.get("default_customer_code") or ""
    threshold = settings.get("match_threshold", 0.5)
    match_log = []

    invoices = []
    by_number = {}   # شماره فاکتور -> فاکتور (ردیف‌های با شماره یکسان، یک فاکتور هستند)

    for row in ws.iter_rows(min_row=2, values_only=True):
        item = row[i_item] if i_item < len(row) else None
        qty = row[i_qty] if i_qty < len(row) else None
        price = row[i_price] if i_price < len(row) else None
        if item in (None, "") and qty in (None, "") and price in (None, ""):
            continue

        inv_no = num_to_str(row[i_inv]) if i_inv < len(row) else ""
        current = by_number.get(inv_no) if inv_no != "" else None

        if current is None:
            raw_customer = row[i_cust] if i_cust < len(row) else None
            raw_customer = str(raw_customer).strip() if raw_customer not in (None, "") else ""
            date_val = row[i_date] if i_date < len(row) else None
            code, matched_name, score = find_customer_code(raw_customer, customer_list, threshold, default_code)
            match_log.append((raw_customer, matched_name, score, code))
            current = {"invoice_no": inv_no, "customer_raw": raw_customer, "customer_code": code, "date": date_val, "items": []}
            invoices.append(current)
            if inv_no != "":
                by_number[inv_no] = current

        current["items"].append({
            "code": num_to_str(item),
            "qty": num_to_str(qty),
            "price": num_to_str(price),
        })

    return invoices, match_log


# ---------------------------------------------------------------------
# اجرای واقعی
# ---------------------------------------------------------------------
def run_all_invoices(invoices, window_title, log_widget):
    def log(msg):
        log_widget.insert(tk.END, msg + "\n")
        log_widget.see(tk.END)
        log_widget.update()

    abort_flag["stop"] = False
    listener = start_panic_listener()

    log(f"شروع ثبت {len(invoices)} فاکتور در ۵ ثانیه... (برای توقف هر لحظه Esc را بزنید)")
    for i in range(5, 0, -1):
        log(f"  {i}...")
        time.sleep(1)

    try:
        activate_window(window_title)
        for n, inv in enumerate(invoices, 1):
            check_abort()
            log(f"[{n}/{len(invoices)}] در حال ثبت فاکتور «{inv['customer_raw']}» (کد {inv['customer_code']}) ...")
            write_invoice(inv["customer_code"], inv["date"], inv["items"], first_invoice=(n == 1), log_func=log)
            time.sleep(scaled(0.4))
        log("پایان: همه‌ی فاکتورها ارسال شدند. لطفاً داخل محک بررسی کنید.")
    except KeyboardInterrupt:
        log("⛔ متوقف شد توسط کاربر (Esc). فاکتورهای بعد از این نقطه ثبت نشدند.")
    except Exception as e:
        log(f"❌ خطا: {e}")
    finally:
        listener.stop()


# ---------------------------------------------------------------------
# رابط گرافیکی
# ---------------------------------------------------------------------
def step_map_columns():
    path = filedialog.askopenfilename(title="یک نمونه از فایل فروش خام انتخاب کنید", filetypes=[("Excel files", "*.xlsx *.xls")])
    if not path:
        return
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.active
    headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1)) if c.value not in (None, "")]
    if not headers:
        messagebox.showerror("خطا", "ستون هدر در سطر اول پیدا نشد.")
        return

    top = tk.Toplevel()
    top.title("تطبیق ستون‌ها")
    top.geometry("420x460")
    tk.Label(top, text="ستون متناظر هر فیلد را انتخاب کنید:", wraplength=380).pack(pady=8)

    fields = [
        ("invoice_no", "شماره فاکتور (ردیف‌های با شماره یکسان = یک فاکتور با چند قلم)"),
        ("customer", "نام مشتری (متن، برای تطبیق با لیست اشخاص)"),
        ("date", "تاریخ فروش (متن شمسی، مثل 1405/07/04)"),
        ("item", "کد کالا (همون عدد یکتایی که در محک ثبت شده)"),
        ("qty", "تعداد/مقدار"),
        ("price", "قیمت واحد"),
    ]
    settings = load_settings()
    current = settings["columns"]
    vars_ = {}
    for key, label in fields:
        tk.Label(top, text=label, wraplength=380, justify="right").pack(anchor="e", padx=10)
        v = tk.StringVar(value=current.get(key, ""))
        cb = ttk.Combobox(top, textvariable=v, values=headers, width=48, state="readonly")
        cb.pack(padx=10, pady=(0, 6))
        vars_[key] = v

    def save():
        if any(not v.get() for v in vars_.values()):
            messagebox.showerror("خطا", "لطفاً همه‌ی فیلدها را مشخص کنید.")
            return
        s = load_settings()
        s["columns"] = {k: v.get() for k, v in vars_.items()}
        s["columns_set"] = True
        save_settings(s)
        messagebox.showinfo("ذخیره شد", "تنظیمات ستون‌ها ذخیره شد.")
        top.destroy()

    tk.Button(top, text="ذخیره", command=save, width=20).pack(pady=10)


def step_pick_window():
    settings = load_settings()
    proceed = messagebox.askokcancel(
        "انتخاب پنجره",
        "لطفاً مطمئن شوید نرم‌افزار محک الان باز است، سپس OK بزنید تا لیست پنجره‌های باز نمایش داده شود.",
    )
    if not proceed:
        return

    titles = sorted(set(t for t in gw.getAllTitles() if t.strip()))
    if not titles:
        messagebox.showerror("خطا", "هیچ پنجره‌ای پیدا نشد.")
        return

    top = tk.Toplevel()
    top.title("انتخاب پنجره‌ی محک")
    top.geometry("500x400")
    tk.Label(top, text="پنجره‌ی نرم‌افزار محک را از لیست زیر انتخاب کنید:").pack(pady=8)

    listbox = tk.Listbox(top, width=70, height=18)
    for t in titles:
        listbox.insert(tk.END, t)
    listbox.pack(padx=10, pady=5)

    def save():
        sel = listbox.curselection()
        if not sel:
            messagebox.showerror("خطا", "یک پنجره را انتخاب کنید.")
            return
        title = listbox.get(sel[0])
        s = load_settings()
        s["window_title"] = title
        s["window_set"] = True
        save_settings(s)
        messagebox.showinfo("ذخیره شد", f"پنجره‌ی زیر ذخیره شد:\n{title}")
        top.destroy()

    tk.Button(top, text="ذخیره", command=save, width=20).pack(pady=10)


def step_run(test_mode_var, log_widget):
    settings = load_settings()
    if not settings.get("columns_set"):
        messagebox.showerror("خطا", "ابتدا ستون‌ها را تنظیم کنید.")
        return
    if not test_mode_var.get() and not settings.get("window_set"):
        messagebox.showerror("خطا", "ابتدا پنجره‌ی محک را انتخاب کنید (یا حالت آزمایشی را فعال کنید).")
        return
    if not settings.get("default_customer_code"):
        messagebox.showerror("خطا", "ابتدا «کد مشتری پیش‌فرض» (مثلاً کد مشتری حضوری) را در پایین صفحه وارد و ذخیره کنید.")
        return

    customers_path = filedialog.askopenfilename(
        title="فایل «لیست اشخاص» همین شرکت را انتخاب کنید (کد + نام)",
        filetypes=[("Excel files", "*.xlsx *.xls")],
    )
    if not customers_path:
        return
    try:
        customer_list = load_customer_lookup(customers_path)
    except Exception as e:
        messagebox.showerror("خطا", str(e))
        return

    if customer_list is None:
        # عنوان ستون پیدا نشد؛ کاربر باید دستی ستون‌ها را مشخص کند
        def continue_with(cl):
            _continue_run_after_customers(test_mode_var, log_widget, settings, cl)

        pick_customer_columns_dialog(customers_path, continue_with)
        return

    if not customer_list:
        messagebox.showwarning("توجه", "لیست اشخاص خالی به نظر می‌رسد.")
        return

    _continue_run_after_customers(test_mode_var, log_widget, settings, customer_list)


def _continue_run_after_customers(test_mode_var, log_widget, settings, customer_list):
    raw_path = filedialog.askopenfilename(title="فایل فروش این دوره را انتخاب کنید", filetypes=[("Excel files", "*.xlsx *.xls")])
    if not raw_path:
        return

    try:
        invoices, match_log = build_invoices(raw_path, settings, customer_list)
    except Exception as e:
        messagebox.showerror("خطا", str(e))
        return

    if not invoices:
        messagebox.showwarning("توجه", "هیچ ردیف قابل‌استفاده‌ای پیدا نشد.")
        return

    # نمایش خلاصه‌ی تطبیق مشتری‌ها برای بررسی قبل از اجرای واقعی
    review = tk.Toplevel()
    review.title("بررسی تطبیق مشتری‌ها")
    review.geometry("640x420")
    tk.Label(review, text="لطفاً قبل از ادامه، تطبیق‌ها را بررسی کنید:", font=("Tahoma", 10, "bold")).pack(pady=8)
    text = tk.Text(review, width=90, height=20)
    text.pack(padx=10, pady=5)
    for raw_name, matched_name, score, code in match_log:
        if matched_name:
            text.insert(tk.END, f"«{raw_name}»  →  «{matched_name}» (کد {code})  —  شباهت {score*100:.0f}٪\n")
        else:
            text.insert(tk.END, f"«{raw_name}»  →  پیدا نشد، کد پیش‌فرض «{code}» استفاده می‌شود  (بیشترین شباهت {score*100:.0f}٪)\n")
    text.config(state="disabled")

    result = {"proceed": False}

    def do_proceed():
        result["proceed"] = True
        review.destroy()

    def do_cancel():
        review.destroy()

    btn_frame = tk.Frame(review)
    btn_frame.pack(pady=8)
    tk.Button(btn_frame, text="تأیید و ادامه", command=do_proceed, width=20, bg="#5cb85c", fg="white").pack(side="left", padx=5)
    tk.Button(btn_frame, text="لغو", command=do_cancel, width=20).pack(side="left", padx=5)
    review.grab_set()
    review.wait_window()

    if not result["proceed"]:
        return

    total_items = sum(len(i["items"]) for i in invoices)
    proceed = messagebox.askokcancel(
        "تأیید نهایی",
        f"{len(invoices)} فاکتور شامل {total_items} ردیف کالا آماده‌ی ثبت است.\n\n"
        + ("حالت آزمایشی فعال است: به‌جای محک در Notepad تایپ می‌شود.\n\n" if test_mode_var.get() else "⚠️ حالت واقعی: مستقیماً داخل محک ثبت می‌شود.\n\n")
        + "ادامه می‌دهید؟",
    )
    if not proceed:
        return

    if test_mode_var.get():
        subprocess.Popen(["notepad.exe"])
        time.sleep(1.2)
        window_title = "Notepad"
    else:
        window_title = settings["window_title"]

    threading.Thread(target=run_all_invoices, args=(invoices, window_title, log_widget), daemon=True).start()


def main():
    settings = load_settings()
    root = tk.Tk()
    root.title("نوشتن خودکار فاکتور در محک")
    root.geometry("560x520")

    tk.Label(root, text="اتوماسیون ثبت فاکتور فروش در محک", font=("Tahoma", 12, "bold")).pack(pady=10)

    tk.Button(root, text="۱) تنظیم ستون‌های فایل خام", command=step_map_columns, width=50).pack(pady=5)
    tk.Button(root, text="۲) انتخاب پنجره‌ی محک از لیست پنجره‌های باز", command=step_pick_window, width=50).pack(pady=5)

    test_mode_var = tk.BooleanVar(value=True)
    tk.Checkbutton(root, text="حالت آزمایشی (به‌جای محک، در Notepad تایپ کن)", variable=test_mode_var).pack(pady=8)

    tk.Label(root, text="ضریب کندی (عدد بزرگ‌تر = آهسته‌تر و مطمئن‌تر):").pack()
    speed_var = tk.DoubleVar(value=settings.get("speed_multiplier", 3.0))
    speed_scale = tk.Scale(root, from_=1.0, to=8.0, resolution=0.5, orient="horizontal", length=300, variable=speed_var)
    speed_scale.pack()

    def apply_speed(*_):
        SPEED["value"] = speed_var.get()
        s = load_settings()
        s["speed_multiplier"] = speed_var.get()
        save_settings(s)

    speed_scale.bind("<ButtonRelease-1>", apply_speed)
    SPEED["value"] = speed_var.get()

    log_widget = tk.Text(root, width=68, height=16)
    log_widget.pack(pady=8)

    tk.Button(root, text="۳) اجرا: تبدیل و ثبت فاکتورهای این دوره", command=lambda: step_run(test_mode_var, log_widget), width=50, bg="#d9534f", fg="white").pack(pady=8)

    tk.Label(root, text="کد مشتری پیش‌فرض برای موارد ناشناس (مثلاً کد «مشتری حضوری»):").pack(pady=(10, 0))
    default_var = tk.StringVar(value=settings.get("default_customer_code", ""))
    tk.Entry(root, textvariable=default_var, width=40, justify="right").pack()

    tk.Label(root, text="حداقل درصد شباهت برای پذیرفتن تطابق مشتری (۰ تا ۱):").pack(pady=(8, 0))
    threshold_var = tk.DoubleVar(value=settings.get("match_threshold", 0.5))
    tk.Scale(root, from_=0.1, to=1.0, resolution=0.05, orient="horizontal", length=250, variable=threshold_var).pack()

    def save_default():
        s = load_settings()
        s["default_customer_code"] = default_var.get().strip()
        s["match_threshold"] = threshold_var.get()
        save_settings(s)
        messagebox.showinfo("ذخیره شد", "ذخیره شد.")

    tk.Button(root, text="ذخیره", command=save_default, width=20).pack(pady=6)

    root.mainloop()


if __name__ == "__main__":
    main()
