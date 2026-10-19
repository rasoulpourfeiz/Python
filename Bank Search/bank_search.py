# -*- coding: utf-8 -*-
"""
Excel Search App
جستجو در همه شیت‌های یک فایل اکسل و ذخیره نتایج در فایل خروجی
"""

import os
import re
import sys
import subprocess
import threading
import tkinter as tk
from tkinter import ttk, filedialog, messagebox

import pandas as pd
import openpyxl


# ==== تبدیل اعداد فارسی و عربی به انگلیسی ====
PERSIAN_NUMS = "۰۱۲۳۴۵۶۷۸۹"
ARABIC_NUMS = "٠١٢٣٤٥٦٧٨٩"
ENGLISH_NUMS = "0123456789"
TRANS_TABLE = str.maketrans(PERSIAN_NUMS + ARABIC_NUMS, ENGLISH_NUMS * 2)


def persian_to_english(s):
    if s is None:
        return s
    try:
        if pd.isna(s):
            return s
    except (TypeError, ValueError):
        pass
    return str(s).translate(TRANS_TABLE)


def search_workbook(input_file, query, output_file, log=print, progress=None):
    """
    جستجو در همه شیت‌ها. خروجی: تعداد کل ردیف‌های پیدا شده.
    """
    query = persian_to_english(query).strip()
    pattern = re.compile(re.escape(query), re.IGNORECASE)

    wb = openpyxl.load_workbook(input_file, data_only=True, read_only=True)
    sheet_names = wb.sheetnames
    total_sheets = len(sheet_names)

    matched_sheets = []
    total_rows = 0

    for idx, sheet_name in enumerate(sheet_names, start=1):
        ws = wb[sheet_name]
        rows = ws.iter_rows(values_only=True)

        try:
            header = next(rows)
        except StopIteration:
            log(f"⚠️ شیت «{sheet_name}» خالی است، رد شد.")
            if progress:
                progress(idx, total_sheets)
            continue

        # نام ستون‌ها را رشته می‌کنیم تا ستون بدون نام مشکل ایجاد نکند
        cols = [
            str(c) if c is not None else f"Column_{i + 1}"
            for i, c in enumerate(header)
        ]

        matched_rows = []
        for row in rows:
            for val in row:
                if val is None:
                    continue
                if pattern.search(persian_to_english(val)):
                    matched_rows.append(row)
                    break

        if matched_rows:
            # اگر ردیف‌ها طول متفاوت داشتند هم‌اندازه می‌کنیم
            width = len(cols)
            fixed = [list(r[:width]) + [None] * (width - len(r)) for r in matched_rows]
            df = pd.DataFrame(fixed, columns=cols)
            df["Sheet_Name"] = sheet_name
            matched_sheets.append((sheet_name, df))
            total_rows += len(df)
            log(f"✅ شیت «{sheet_name}»: {len(df)} ردیف پیدا شد")
        else:
            log(f"— شیت «{sheet_name}»: نتیجه‌ای نبود")

        if progress:
            progress(idx, total_sheets)

    wb.close()

    if matched_sheets:
        out_dir = os.path.dirname(output_file)
        if out_dir:
            os.makedirs(out_dir, exist_ok=True)
        with pd.ExcelWriter(output_file, engine="openpyxl") as writer:
            used_names = set()
            for sheet_name, df in matched_sheets:
                # محدودیت اکسل: حداکثر ۳۱ کاراکتر و بدون کاراکترهای ممنوع
                safe = re.sub(r"[\\/*?:\[\]]", "_", sheet_name)[:31] or "Sheet"
                base, n = safe, 1
                while safe in used_names:
                    suffix = f"_{n}"
                    safe = base[: 31 - len(suffix)] + suffix
                    n += 1
                used_names.add(safe)
                df.to_excel(writer, index=False, sheet_name=safe)

    return total_rows, bool(matched_sheets)


# ==================== رابط گرافیکی ====================
class App(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("جستجو در فایل اکسل | Excel Search")
        self.geometry("720x560")
        self.minsize(640, 480)

        self.input_var = tk.StringVar()
        self.output_var = tk.StringVar()
        self.query_var = tk.StringVar()
        self.last_output = None

        self._build_ui()

    def _build_ui(self):
        pad = {"padx": 10, "pady": 6}
        frm = ttk.Frame(self)
        frm.pack(fill="both", expand=True, padx=10, pady=10)
        frm.columnconfigure(1, weight=1)
        frm.rowconfigure(5, weight=1)

        # فایل ورودی
        ttk.Label(frm, text="فایل ورودی (Excel):").grid(row=0, column=0, sticky="w", **pad)
        ttk.Entry(frm, textvariable=self.input_var).grid(row=0, column=1, sticky="ew", **pad)
        ttk.Button(frm, text="انتخاب...", command=self.browse_input).grid(row=0, column=2, **pad)

        # عبارت جستجو
        ttk.Label(frm, text="عبارت جستجو:").grid(row=1, column=0, sticky="w", **pad)
        entry = ttk.Entry(frm, textvariable=self.query_var)
        entry.grid(row=1, column=1, columnspan=2, sticky="ew", **pad)
        entry.bind("<Return>", lambda e: self.start_search())
        entry.focus_set()

        # فایل خروجی
        ttk.Label(frm, text="فایل خروجی:").grid(row=2, column=0, sticky="w", **pad)
        ttk.Entry(frm, textvariable=self.output_var).grid(row=2, column=1, sticky="ew", **pad)
        ttk.Button(frm, text="ذخیره در...", command=self.browse_output).grid(row=2, column=2, **pad)

        # دکمه‌ها
        btns = ttk.Frame(frm)
        btns.grid(row=3, column=0, columnspan=3, sticky="ew", **pad)
        self.search_btn = ttk.Button(btns, text="🔍 جستجو", command=self.start_search)
        self.search_btn.pack(side="left")
        self.open_btn = ttk.Button(btns, text="📂 باز کردن پوشه خروجی",
                                   command=self.open_output_folder, state="disabled")
        self.open_btn.pack(side="left", padx=8)

        # نوار پیشرفت
        self.progress = ttk.Progressbar(frm, mode="determinate")
        self.progress.grid(row=4, column=0, columnspan=3, sticky="ew", **pad)

        # گزارش
        self.log_box = tk.Text(frm, height=10, wrap="word", state="disabled")
        self.log_box.grid(row=5, column=0, columnspan=3, sticky="nsew", **pad)
        scroll = ttk.Scrollbar(frm, command=self.log_box.yview)
        scroll.grid(row=5, column=3, sticky="ns")
        self.log_box.config(yscrollcommand=scroll.set)

        # مسیر خروجی نهایی
        self.result_label = ttk.Label(frm, text="", foreground="#1a7f37", wraplength=660)
        self.result_label.grid(row=6, column=0, columnspan=3, sticky="w", **pad)

    # ---------- کمکی‌ها ----------
    def log(self, msg):
        def _do():
            self.log_box.config(state="normal")
            self.log_box.insert("end", msg + "\n")
            self.log_box.see("end")
            self.log_box.config(state="disabled")
        self.after(0, _do)

    def set_progress(self, done, total):
        self.after(0, lambda: self.progress.config(maximum=total, value=done))

    def browse_input(self):
        path = filedialog.askopenfilename(
            title="انتخاب فایل اکسل",
            filetypes=[("Excel files", "*.xlsx *.xlsm"), ("All files", "*.*")],
        )
        if path:
            self.input_var.set(path)
            if not self.output_var.get():
                self.output_var.set(
                    os.path.join(os.path.dirname(path), "Search_Results.xlsx")
                )

    def browse_output(self):
        initial = self.output_var.get()
        path = filedialog.asksaveasfilename(
            title="محل ذخیره فایل خروجی",
            defaultextension=".xlsx",
            initialfile=os.path.basename(initial) if initial else "Search_Results.xlsx",
            initialdir=os.path.dirname(initial) if initial else None,
            filetypes=[("Excel files", "*.xlsx")],
        )
        if path:
            self.output_var.set(path)

    def open_output_folder(self):
        if not self.last_output:
            return
        folder = os.path.dirname(os.path.abspath(self.last_output))
        try:
            if sys.platform == "darwin":
                subprocess.Popen(["open", folder])
            elif sys.platform.startswith("win"):
                os.startfile(folder)  # type: ignore[attr-defined]
            else:
                subprocess.Popen(["xdg-open", folder])
        except Exception as e:
            messagebox.showerror("خطا", f"باز کردن پوشه ممکن نشد:\n{e}")

    # ---------- اجرای جستجو ----------
    def start_search(self):
        input_file = self.input_var.get().strip()
        query = self.query_var.get().strip()
        output_file = self.output_var.get().strip()

        if not input_file or not os.path.isfile(input_file):
            messagebox.showwarning("توجه", "لطفاً یک فایل اکسل معتبر انتخاب کنید.")
            return
        if not query:
            messagebox.showwarning("توجه", "لطفاً عبارت جستجو را وارد کنید.")
            return
        if not output_file:
            output_file = os.path.join(os.path.dirname(input_file), "Search_Results.xlsx")
            self.output_var.set(output_file)
        if not output_file.lower().endswith(".xlsx"):
            output_file += ".xlsx"
            self.output_var.set(output_file)
        if os.path.abspath(output_file) == os.path.abspath(input_file):
            messagebox.showwarning("توجه", "فایل خروجی نباید همان فایل ورودی باشد.")
            return

        self.search_btn.config(state="disabled")
        self.open_btn.config(state="disabled")
        self.result_label.config(text="")
        self.progress.config(value=0)
        self.log_box.config(state="normal")
        self.log_box.delete("1.0", "end")
        self.log_box.config(state="disabled")
        self.log("⏳ در حال جستجو...")

        threading.Thread(
            target=self._worker, args=(input_file, query, output_file), daemon=True
        ).start()

    def _worker(self, input_file, query, output_file):
        try:
            total_rows, found = search_workbook(
                input_file, query, output_file,
                log=self.log, progress=self.set_progress,
            )
            self.after(0, lambda: self._finish(found, total_rows, query, output_file))
        except PermissionError:
            self.after(0, lambda: self._error(
                "دسترسی به فایل ممکن نیست. اگر فایل خروجی در اکسل باز است آن را ببندید."))
        except Exception as e:
            self.after(0, lambda: self._error(str(e)))

    def _finish(self, found, total_rows, query, output_file):
        self.search_btn.config(state="normal")
        if found:
            self.last_output = output_file
            self.open_btn.config(state="normal")
            self.result_label.config(
                text=f"✅ {total_rows} ردیف پیدا شد.\nمسیر فایل خروجی:\n{output_file}"
            )
            self.log(f"\n✅ جستجو کامل شد! نتایج در:\n{output_file}")
            messagebox.showinfo("پایان", f"{total_rows} ردیف پیدا شد.\n\nفایل خروجی:\n{output_file}")
        else:
            self.result_label.config(text=f"❌ هیچ داده‌ای برای «{query}» پیدا نشد.",
                                     foreground="#b00020")
            self.log(f"\n❌ هیچ داده‌ای برای «{query}» پیدا نشد (فایل خروجی ساخته نشد).")
            messagebox.showinfo("نتیجه", f"هیچ داده‌ای برای «{query}» پیدا نشد.")

    def _error(self, msg):
        self.search_btn.config(state="normal")
        self.log(f"❌ خطا: {msg}")
        messagebox.showerror("خطا", msg)


if __name__ == "__main__":
    App().mainloop()
