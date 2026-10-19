"""هسته‌ی سیستم محاسبه سود سرمایه‌گذاری: دیتابیس SQLite، محاسبه سود، گزارش‌ها، ورود/خروج اکسل."""
import io
import os
import re
import sqlite3
from collections import defaultdict
from datetime import date, timedelta
from decimal import Decimal, ROUND_HALF_UP

import jalali as J

BASE = os.path.dirname(os.path.abspath(__file__))
DB_PATH = os.environ.get("INVEST_DB") or os.path.join(BASE, "data", "invest.db")

DEPOSIT, WITHDRAW = "واریز", "برداشت"
ACTIVE, INACTIVE = "فعال", "غیرفعال"
DEFAULT_SETTINGS = {"monthly_rate": "0.08", "day_basis": "30", "rounding": "1000"}

INV_FIELDS = ["national_id", "contract_no", "contract_date", "first_name", "last_name",
              "father_name", "birth_date", "mobile", "phone", "province", "city",
              "address", "notes", "status"]

SCHEMA = """
CREATE TABLE IF NOT EXISTS settings(key TEXT PRIMARY KEY, value TEXT NOT NULL);
CREATE TABLE IF NOT EXISTS investors(
  national_id TEXT PRIMARY KEY, contract_no TEXT, contract_date TEXT,
  first_name TEXT, last_name TEXT, father_name TEXT, birth_date TEXT,
  mobile TEXT, phone TEXT, province TEXT, city TEXT, address TEXT, notes TEXT,
  status TEXT NOT NULL DEFAULT 'فعال');
CREATE TABLE IF NOT EXISTS banks(
  national_id TEXT NOT NULL REFERENCES investors(national_id) ON DELETE CASCADE,
  slot INTEGER NOT NULL, bank TEXT, account TEXT, card TEXT, sheba TEXT,
  PRIMARY KEY(national_id, slot));
CREATE TABLE IF NOT EXISTS transactions(
  id INTEGER PRIMARY KEY AUTOINCREMENT,
  national_id TEXT NOT NULL REFERENCES investors(national_id),
  tx_date TEXT NOT NULL, tx_type TEXT NOT NULL CHECK(tx_type IN ('واریز','برداشت')),
  amount INTEGER NOT NULL CHECK(amount > 0), notes TEXT);
CREATE TABLE IF NOT EXISTS payments(
  id INTEGER PRIMARY KEY AUTOINCREMENT,
  national_id TEXT NOT NULL REFERENCES investors(national_id),
  pay_date TEXT NOT NULL, from_date TEXT NOT NULL, to_date TEXT NOT NULL,
  amount INTEGER NOT NULL, calc_amount INTEGER NOT NULL, notes TEXT);
CREATE INDEX IF NOT EXISTS ix_tx ON transactions(national_id, tx_date);
CREATE INDEX IF NOT EXISTS ix_pay ON payments(national_id, to_date);
"""


# ───────────────────────── دیتابیس و تنظیمات ─────────────────────────
def connect(path=None):
    path = path or DB_PATH
    os.makedirs(os.path.dirname(path), exist_ok=True)
    con = sqlite3.connect(path)
    con.row_factory = sqlite3.Row
    con.execute("PRAGMA foreign_keys=ON")
    con.executescript(SCHEMA)
    for k, v in DEFAULT_SETTINGS.items():
        con.execute("INSERT OR IGNORE INTO settings VALUES(?,?)", (k, v))
    con.commit()
    return con


def get_settings(con):
    return {r["key"]: r["value"] for r in con.execute("SELECT * FROM settings")}


def save_settings(con, data):
    try:
        rate = Decimal(J.to_en_digits(data["monthly_rate"]).replace("٫", ".").replace("٪", "").strip())
        basis = int(J.to_en_digits(data["day_basis"]))
        rnd = int(J.to_en_digits(data["rounding"]).replace(",", ""))
    except Exception:
        raise ValueError("مقادیر تنظیمات نامعتبر است")
    if rate > 1:  # کاربر ۸ نوشته نه ۰٫۰۸
        rate = rate / 100
    if not (Decimal(0) < rate < 1) or basis <= 0 or rnd < 0:
        raise ValueError("نرخ باید بین ۰ تا ۱ (مثلاً 0.08 یا 8) و مبنای روز بزرگ‌تر از صفر باشد")
    for k, v in (("monthly_rate", str(rate)), ("day_basis", str(basis)), ("rounding", str(rnd))):
        con.execute("INSERT OR REPLACE INTO settings VALUES(?,?)", (k, v))
    con.commit()


# ───────────────────────── ابزارهای کمکی ─────────────────────────
def clean_id(v):
    return J.to_en_digits(v if v is not None else "").strip()


def to_int(v, what="مبلغ"):
    s = J.to_en_digits(v if v is not None else "")
    s = re.sub(r"[,\s٬،]", "", s)
    if not re.fullmatch(r"\d+", s) or int(s) <= 0:
        raise ValueError(f"{what} باید یک عدد صحیح مثبت (ریال) باشد")
    return int(s)


def iso(d):
    return d.isoformat() if d else None


def from_iso(s):
    return date.fromisoformat(s) if s else None


# ───────────────────────── سرمایه‌گذاران ─────────────────────────
def upsert_investor(con, data, banks=None):
    nid = clean_id(data.get("national_id"))
    if not nid:
        raise ValueError("کد ملی / شناسه مشتری الزامی است")
    if not (str(data.get("first_name") or "").strip() or str(data.get("last_name") or "").strip()):
        raise ValueError("نام یا نام خانوادگی الزامی است")
    status = data.get("status") or ACTIVE
    if status not in (ACTIVE, INACTIVE):
        raise ValueError("وضعیت باید «فعال» یا «غیرفعال» باشد")
    row = {k: (str(data.get(k) or "").strip()) for k in INV_FIELDS}
    row["national_id"], row["status"] = nid, status
    for k in ("mobile", "phone", "contract_no"):
        row[k] = J.to_en_digits(row[k])
    row["contract_date"] = iso(J.parse(row["contract_date"])) if row["contract_date"] else ""
    row["birth_date"] = J.fmt(J.parse(row["birth_date"])) if row["birth_date"] else ""
    cols = ",".join(INV_FIELDS)
    con.execute(f"INSERT OR REPLACE INTO investors({cols}) VALUES({','.join('?'*len(INV_FIELDS))})",
                [row[k] for k in INV_FIELDS])
    if banks is not None:
        con.execute("DELETE FROM banks WHERE national_id=?", (nid,))
        for i, b in enumerate(banks[:4], 1):
            vals = [J.to_en_digits(str(b.get(k) or "").strip()) if k != "bank" else str(b.get(k) or "").strip()
                    for k in ("bank", "account", "card", "sheba")]
            if any(vals):
                con.execute("INSERT INTO banks VALUES(?,?,?,?,?,?)", (nid, i, *vals))
    con.commit()
    return nid


def get_investor(con, nid):
    r = con.execute("SELECT * FROM investors WHERE national_id=?", (clean_id(nid),)).fetchone()
    if not r:
        raise ValueError("سرمایه‌گذار پیدا نشد")
    d = dict(r)
    d["contract_date"] = J.fmt(from_iso(d["contract_date"]))
    banks = {b["slot"]: dict(b) for b in con.execute("SELECT * FROM banks WHERE national_id=?", (nid,))}
    d["banks"] = [{k: (banks.get(i) or {}).get(k, "") or "" for k in ("bank", "account", "card", "sheba")}
                  for i in range(1, 5)]
    return d


def list_investors(con):
    return [dict(r) for r in con.execute(
        "SELECT national_id, first_name, last_name, contract_no, mobile, status "
        "FROM investors ORDER BY last_name, first_name")]


# ───────────────────────── تراکنش‌ها ─────────────────────────
def add_tx(con, nid, date_str, tx_type, amount, notes=""):
    nid = clean_id(nid)
    if not con.execute("SELECT 1 FROM investors WHERE national_id=?", (nid,)).fetchone():
        raise ValueError("سرمایه‌گذار پیدا نشد")
    if tx_type not in (DEPOSIT, WITHDRAW):
        raise ValueError("نوع تراکنش باید واریز یا برداشت باشد")
    d, amt = J.parse(date_str), to_int(amount)
    if tx_type == WITHDRAW:
        bal = con.execute("SELECT COALESCE(SUM(CASE tx_type WHEN 'واریز' THEN amount ELSE -amount END),0) "
                          "FROM transactions WHERE national_id=?", (nid,)).fetchone()[0]
        if amt > bal:
            raise ValueError(f"مبلغ برداشت ({amt:,}) بیشتر از سرمایه فعلی ({bal:,}) است")
    con.execute("INSERT INTO transactions(national_id,tx_date,tx_type,amount,notes) VALUES(?,?,?,?,?)",
                (nid, iso(d), tx_type, amt, str(notes or "").strip()))
    con.commit()


def list_tx(con, nid=None):
    q = ("SELECT t.id, t.national_id, i.first_name, i.last_name, t.tx_date, t.tx_type, t.amount, t.notes "
         "FROM transactions t JOIN investors i USING(national_id) ")
    args = []
    if nid:
        q += "WHERE t.national_id=? "
        args.append(clean_id(nid))
    out = []
    for r in con.execute(q + "ORDER BY t.tx_date DESC, t.id DESC", args):
        d = dict(r)
        d["date_j"] = J.fmt(from_iso(d["tx_date"]))
        out.append(d)
    return out


def delete_tx(con, rid):
    con.execute("DELETE FROM transactions WHERE id=?", (int(rid),))
    con.commit()


def delete_payment(con, pid):
    r = con.execute("SELECT national_id, to_date FROM payments WHERE id=?", (int(pid),)).fetchone()
    if not r:
        raise ValueError("پرداخت پیدا نشد")
    last = con.execute("SELECT MAX(to_date) FROM payments WHERE national_id=?", (r["national_id"],)).fetchone()[0]
    if r["to_date"] != last:
        raise ValueError("فقط آخرین پرداخت هر سرمایه‌گذار قابل حذف است")
    con.execute("DELETE FROM payments WHERE id=?", (int(pid),))
    con.commit()


# ───────────────────────── محاسبه سود ─────────────────────────
def weighted_day_sum(txs, last_pay, as_of):
    """Σ (±مبلغ × روز) برای تراکنش‌های قبل از as_of؛ شروع شمارش = max(تاریخ تراکنش، آخرین پرداخت)"""
    total = 0
    for d, signed in txs:
        if d < as_of:
            days = (as_of - max(d, last_pay)).days
            if days > 0:
                total += signed * days
    return total


def calc_interest(weighted, st):
    raw = Decimal(st["monthly_rate"]) / Decimal(st["day_basis"]) * Decimal(weighted)
    m = Decimal(st["rounding"])
    if m <= 0:
        return int(raw.quantize(Decimal(1), ROUND_HALF_UP))
    q = (abs(raw) / m).quantize(Decimal(1), rounding=ROUND_HALF_UP) * m   # مثل MROUND اکسل
    return int(q if raw >= 0 else -q)


def _load(con):
    txs = defaultdict(list)
    for r in con.execute("SELECT national_id, tx_date, tx_type, amount FROM transactions"):
        txs[r["national_id"]].append((from_iso(r["tx_date"]), r["amount"] if r["tx_type"] == DEPOSIT else -r["amount"]))
    paid = {r[0]: from_iso(r[1]) for r in con.execute("SELECT national_id, MAX(to_date) FROM payments GROUP BY 1")}
    return txs, paid


def next_period(first, paid_through):
    """دوره‌ی پرداخت‌نشده‌ی بعدی -> (شروع، پایان=سررسید، شماره‌ی دوره).
    سررسیدها همیشه سالگرد ماهانه‌ی «اولین تاریخ سرمایه‌گذاری» هستند (به تقویم شمسی)،
    صرف‌نظر از اینکه سود کی واقعاً پرداخت شده باشد."""
    k = 1
    while paid_through and J.add_months(first, k) <= paid_through:
        k += 1
    return (paid_through or first), J.add_months(first, k), k


def _state(inv, txs, paid, today, st):
    deposits = [d for d, s in txs if s > 0]
    first = min(deposits) if deposits else None
    active = inv["status"] == ACTIVE
    start = due = None
    accrued = interest = overdue_n = 0
    if first:
        start, due, k = next_period(first, paid)
        while J.add_months(first, k + overdue_n) < today:
            overdue_n += 1
        if active:
            accrued = calc_interest(weighted_day_sum(txs, start, today), st)
            interest = calc_interest(weighted_day_sum(txs, start, due), st)
    if not first:
        text = ""
    elif not active:
        text = "سود تعلق نمی‌گیرد"
    else:
        n = (due - today).days
        text = (f"سررسید گذشته ({-n} روز)" + (f" — {overdue_n} دوره" if overdue_n > 1 else "") if n < 0
                else "امروز سررسید است" if n == 0 else "فردا سررسید است" if n == 1 else f"در انتظار سررسید ({n} روز)")
    return {
        "national_id": inv["national_id"], "first_name": inv["first_name"], "last_name": inv["last_name"],
        "contract_no": inv["contract_no"], "status": inv["status"], "active": active,
        "first_date": J.fmt(first), "capital": sum(s for _, s in txs),
        "last_pay": J.fmt(paid), "due": due, "due_j": J.fmt(due),
        "days_to_due": (due - today).days if due else None, "status_text": text,
        "accrued_today": accrued, "interest_due": interest, "unpaid_periods": overdue_n,
    }


def all_states(con, today=None):
    today, st = today or date.today(), get_settings(con)
    txs, paid = _load(con)
    return [_state(i, txs[i["national_id"]], paid.get(i["national_id"]), today, st)
            for i in con.execute("SELECT * FROM investors ORDER BY last_name, first_name")]


def reports(con, today=None):
    today = today or date.today()
    states = all_states(con, today)
    bank1 = {}
    for b in con.execute("SELECT * FROM banks ORDER BY slot"):
        if (b["bank"] or b["account"] or b["sheba"]) and b["national_id"] not in bank1:
            bank1[b["national_id"]] = b

    def pick(cond):
        rows = []
        for s in sorted((s for s in states if s["active"] and s["due"] and s["interest_due"] > 0 and cond(s["due"])),
                        key=lambda s: (s["due"], s["last_name"] or "")):
            b = bank1.get(s["national_id"])
            rows.append({**s, "due": s["due_j"], "bank": b["bank"] if b else "",
                         "account": b["account"] if b else "", "sheba": b["sheba"] if b else ""})
        return rows

    return {"today": pick(lambda d: d == today),
            "tomorrow": pick(lambda d: d == today + timedelta(days=1)),
            "overdue": pick(lambda d: d < today)}


# ───────────────────────── پرداخت سود ─────────────────────────
def suggest_payment(con, nid, date_str):
    """دوره‌ی بعدیِ پرداخت‌نشده و مبلغ سود آن. تاریخ ورودی = تاریخ واقعی پرداخت (فقط ثبت می‌شود؛
    پایان دوره همیشه سالگرد اولین سرمایه‌گذاری است)."""
    nid, st = clean_id(nid), get_settings(con)
    pay = J.parse(date_str)
    if not con.execute("SELECT 1 FROM investors WHERE national_id=?", (nid,)).fetchone():
        raise ValueError("سرمایه‌گذار پیدا نشد")
    txs, paid = _load(con)
    deposits = [d for d, s in txs[nid] if s > 0]
    if not deposits:
        raise ValueError("برای این سرمایه‌گذار هنوز واریزی ثبت نشده است")
    start, due, k = next_period(min(deposits), paid.get(nid))
    return {"from": J.fmt(start), "to": J.fmt(due), "days": (due - start).days, "period_no": k,
            "amount": calc_interest(weighted_day_sum(txs[nid], start, due), st),
            "_from": start, "_to": due, "_pay": pay}


def add_payment(con, nid, date_str, amount=None, notes=""):
    nid = clean_id(nid)
    s = suggest_payment(con, nid, date_str)
    amt = s["amount"] if amount in (None, "") else to_int(amount, "مبلغ پرداخت")
    con.execute("INSERT INTO payments(national_id,pay_date,from_date,to_date,amount,calc_amount,notes) VALUES(?,?,?,?,?,?,?)",
                (nid, iso(s["_pay"]), iso(s["_from"]), iso(s["_to"]), amt, s["amount"], str(notes or "").strip()))
    con.commit()


def list_payments(con):
    out = []
    for r in con.execute("SELECT p.*, i.first_name, i.last_name FROM payments p JOIN investors i USING(national_id) "
                         "ORDER BY p.pay_date DESC, p.id DESC"):
        d = dict(r)
        d["pay_j"], d["from_j"], d["to_j"] = (J.fmt(from_iso(d[k])) for k in ("pay_date", "from_date", "to_date"))
        out.append(d)
    return out


# ───────────────────────── ورود از اکسل قدیمی ─────────────────────────
def import_excel(con, path):
    """از فایل اکسل نسخه‌ی فعلی (.xlsm / .xlsx): دیتابیس، تراکنش‌ها، پرداخت سود، تنظیمات."""
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True)
    rep = {"investors": 0, "transactions": 0, "payments": 0, "warnings": []}

    def cell(row, i):
        v = row[i] if i < len(row) else None
        return "" if v is None else v

    def text(v):
        if isinstance(v, float) and v.is_integer():
            v = int(v)
        return str(v).strip()

    if "تنظیمات" in wb.sheetnames:
        ws = wb["تنظیمات"]
        try:
            save_settings(con, {"monthly_rate": text(ws["B4"].value), "day_basis": text(ws["B5"].value),
                                "rounding": text(ws["B6"].value)})
        except ValueError:
            rep["warnings"].append("تنظیمات از اکسل خوانده نشد؛ مقادیر پیش‌فرض باقی ماند")

    status = {}
    if "جدول محاسبه سود" in wb.sheetnames:
        for row in wb["جدول محاسبه سود"].iter_rows(min_row=2, values_only=True):
            if row and row[0] is not None:
                status[clean_id(row[0])] = text(cell(row, 5)) or ACTIVE

    ws = wb["دیتابیس"]
    for n, row in enumerate(ws.iter_rows(min_row=2, values_only=True), 2):
        if not row or row[0] in (None, ""):
            continue
        nid = clean_id(text(row[0]))
        if not (text(cell(row, 3)) or text(cell(row, 4))):
            rep["warnings"].append(f"دیتابیس ردیف {n}: شناسه «{nid}» نام ندارد؛ نادیده گرفته شد")
            continue
        data = dict(zip(INV_FIELDS[:13], [nid, *[text(cell(row, i)) for i in range(1, 13)]]))
        data["status"] = status.get(nid, ACTIVE)
        banks = [{"bank": text(cell(row, 13 + 4 * k)), "account": text(cell(row, 14 + 4 * k)),
                  "card": text(cell(row, 15 + 4 * k)), "sheba": text(cell(row, 16 + 4 * k))} for k in range(4)]
        try:
            upsert_investor(con, data, banks)
            rep["investors"] += 1
        except ValueError as e:
            rep["warnings"].append(f"دیتابیس ردیف {n}: {e}")

    known = {r[0] for r in con.execute("SELECT national_id FROM investors")}
    ws = wb["تراکنشهای واریز و برداشت"]
    for n, row in enumerate(ws.iter_rows(min_row=2, values_only=True), 2):
        if not row or row[0] in (None, ""):
            continue
        nid = clean_id(text(row[0]))
        try:
            if nid not in known:
                raise ValueError("سرمایه‌گذار در دیتابیس نیست")
            d, amt = J.parse(text(cell(row, 3))), to_int(cell(row, 5))
            if text(cell(row, 4)) not in (DEPOSIT, WITHDRAW):
                raise ValueError("نوع تراکنش نامعتبر")
            con.execute("INSERT INTO transactions(national_id,tx_date,tx_type,amount,notes) VALUES(?,?,?,?,?)",
                        (nid, iso(d), text(cell(row, 4)), amt, text(cell(row, 6))))
            rep["transactions"] += 1
        except ValueError as e:
            rep["warnings"].append(f"تراکنش ردیف {n}: {e}")
    con.commit()

    if "پرداخت سود" in wb.sheetnames:
        pays = []
        for n, row in enumerate(wb["پرداخت سود"].iter_rows(min_row=2, max_col=8, values_only=True), 2):
            if row and row[0] not in (None, "") and cell(row, 3) != "":
                pays.append((J.parse(text(cell(row, 3))), n, clean_id(text(row[0])), text(cell(row, 7))))
        for d, n, nid, note in sorted(pays):
            try:
                add_payment(con, nid, J.fmt(d), None, note)
                rep["payments"] += 1
            except ValueError as e:
                rep["warnings"].append(f"پرداخت سود ردیف {n}: {e}")
    return rep


# ───────────────────────── خروجی اکسل ─────────────────────────
def make_xlsx(sheets):
    """sheets: {نام شیت: (سرستون‌ها, ردیف‌ها)} -> bytes"""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    wb = Workbook()
    wb.remove(wb.active)
    for name, (headers, rows) in sheets.items():
        ws = wb.create_sheet(name[:31])
        ws.sheet_view.rightToLeft = True
        ws.append(headers)
        for c in ws[1]:
            c.font, c.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1F4E78")
            c.alignment = Alignment(horizontal="center")
        for r in rows:
            ws.append(list(r))
        for i, h in enumerate(headers, 1):
            width = max([len(str(h))] + [len(str(r[i - 1])) for r in rows if r[i - 1] is not None]) + 3
            ws.column_dimensions[ws.cell(1, i).column_letter].width = min(width, 45)
            for row in ws.iter_rows(min_row=2, min_col=i, max_col=i):
                for c in row:
                    if isinstance(c.value, int) and not isinstance(c.value, bool) and abs(c.value) > 99999:
                        c.number_format = "#,##0"
        ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


REPORT_TITLES = {"today": "سررسید امروز", "tomorrow": "سررسید فردا", "overdue": "سررسید گذشته"}
_REPORT_COLS = [("national_id", "شناسه مشتری"), ("first_name", "نام"), ("last_name", "نام خانوادگی"),
                ("bank", "بانک"), ("account", "شماره حساب"), ("sheba", "شماره شبا"),
                ("due", "تاریخ سررسید"), ("interest_due", "مبلغ سود قابل پرداخت (ریال)"),
                ("unpaid_periods", "تعداد دوره پرداخت‌نشده")]


def export_report(con, name, today=None):
    rows = reports(con, today)[name]
    return make_xlsx({REPORT_TITLES[name]: ([h for _, h in _REPORT_COLS],
                                            [[r[k] for k, _ in _REPORT_COLS] for r in rows])})


def export_backup(con, today=None):
    inv = []
    for r in list_investors(con):
        d = get_investor(con, r["national_id"])
        row = [d[k] for k in INV_FIELDS]
        for b in d["banks"]:
            row += [b["bank"], b["account"], b["card"], b["sheba"]]
        inv.append(row)
    ih = ["کدملی", "شماره قرارداد", "تاریخ قرارداد", "نام", "نام خانوادگی", "نام پدر", "تاریخ تولد", "موبایل",
          "تلفن", "استان", "شهر", "آدرس", "توضیحات", "وضعیت"]
    for i in range(1, 5):
        ih += [f"نام بانک {i}", f"شماره حساب {i}", f"شماره کارت {i}", f"شماره شبا {i}"]
    return make_xlsx({
        "سرمایه‌گذاران": (ih, inv),
        "تراکنش‌ها": (["کد ملی", "نام", "نام خانوادگی", "تاریخ", "نوع", "مبلغ (ریال)", "توضیحات"],
                      [[t["national_id"], t["first_name"], t["last_name"], t["date_j"], t["tx_type"], t["amount"], t["notes"]]
                       for t in reversed(list_tx(con))]),
        "پرداخت سود": (["کد ملی", "نام", "نام خانوادگی", "تاریخ پرداخت", "از تاریخ", "تا تاریخ", "مبلغ پرداختی", "توضیحات"],
                       [[p["national_id"], p["first_name"], p["last_name"], p["pay_j"], p["from_j"], p["to_j"], p["amount"], p["notes"]]
                        for p in reversed(list_payments(con))]),
        "جدول محاسبه سود": (["کد ملی", "نام", "نام خانوادگی", "وضعیت", "اولین واریز", "سرمایه فعلی", "آخرین پرداخت",
                             "سررسید بعدی", "سود تا امروز", "سود در سررسید", "وضعیت امروز"],
                            [[s["national_id"], s["first_name"], s["last_name"], s["status"], s["first_date"], s["capital"],
                              s["last_pay"], s["due_j"], s["accrued_today"], s["interest_due"], s["status_text"]]
                             for s in all_states(con, today)]),
    })
